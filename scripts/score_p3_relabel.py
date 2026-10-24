#!/usr/bin/env python3
"""
score_p3_relabel.py — score the blinded relabelling of the P3 page subset and re-run the main
contrasts on the verified pages.

Plan: papers/P3_multimodal/REVISION2/RELABEL_PLAN.md. This file's SHA-256 is recorded there before
any rating exists; do not edit it after the sheets go out (a fix needs a dated amendment in the plan).

Inputs: data/docs/verify_p3/key.csv and the two returned sheets
        data/docs/verify_p3/returned/cham_trang_A.xlsx, cham_trang_B.xlsx
RUN:    python scripts/score_p3_relabel.py [--no-rerun]
OUT:    data/processed/p3/results/relabel.json
        papers/P3_multimodal/sections/gen_relabel.tex, tab_relabel.tex
"""
import argparse
import csv
import os
import random
import sys

import numpy as np

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(_HERE))
try:
    from _path import ROOT, add_script_dirs
    add_script_dirs()
except ImportError:  # flat public-mirror layout
    ROOT = os.path.dirname(_HERE)
from genfile import write_generated, write_results
from paired_eval import bh_adjust, corrected_paired_t, fmt_p as _p

VDIR = os.path.join(ROOT, "data", "docs", "verify_p3")
SEC = os.path.join(ROOT, "papers", "P3_multimodal", "sections")
VERDICTS = ("phishing", "scam", "legitimate", "unsure")
ABUSE = {"phishing", "scam"}
ENCODERS = ("tfidf", "phobert", "phobert-v2", "visobert", "xlm-r")
SPLITS = 20
MLAB = {"F1": "F1", "PR-AUC": "PR-AUC", "ROC-AUC": "ROC-AUC", "FPR@R0.90": "FPR@90\\%rec."}


def collapse(v):
    return "abuse" if v in ABUSE else v


def read_sheet(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True).active
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        item = str(row[0]).strip()
        verdict = str(row[2] or "").strip().lower()
        evidence = str(row[3] or "").strip().lower()
        out[item] = (verdict, evidence)
    return out


def kappa(a, b, cats):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(c) / n) * (b.count(c) / n) for c in cats)
    k = (po - pe) / (1 - pe) if pe < 1 else float("nan")
    pabak = (len(cats) * po - 1) / (len(cats) - 1)
    return po, k, pabak


def wilson(x, n, z=1.96):
    if n == 0:
        return float("nan"), float("nan"), float("nan")
    p = x / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * np.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return p, c - h, c + h


def rerun(verified_rows):
    """The paper's protocol on the verified subset: 20 stratified 70/30 splits, standardised
    linear head, five encoders, both directions, four metrics, corrected t, BH per column."""
    from train_content_fusion import evaluate
    from make_p3_assets import SWEEP_METRICS, SWEEP_DIRECTIONS
    from scipy import stats
    scores = {e: evaluate(verified_rows, ["url", "content", "content+url"], e, "lightweight",
                          False, SPLITS, standardise=True) for e in ENCODERS}
    res = {}
    for d, (a, b) in SWEEP_DIRECTIONS:
        for m, _lab, _s in SWEEP_METRICS:
            tests = [corrected_paired_t(np.asarray(scores[e][a][m], float)
                                        - np.asarray(scores[e][b][m], float)) for e in ENCODERS]
            q, _ = bh_adjust([t["p"] for t in tests])
            for e, t, qi in zip(ENCODERS, tests, q):
                res[(e, d, m)] = {**t, "q": qi}
    mde = {}
    for e in ENCODERS:
        for d, (a, b) in SWEEP_DIRECTIONS:
            dd = np.asarray(scores[e][a]["F1"], float) - np.asarray(scores[e][b]["F1"], float)
            k = len(dd)
            se = np.sqrt((1 / k + 0.3 / 0.7) * np.var(dd, ddof=1))
            mde[(e, d)] = float((stats.t.ppf(0.975, k - 1) + stats.t.ppf(0.80, k - 1)) * se)
    means = {e: {c: float(np.mean(scores[e][c]["F1"])) for c in ("url", "content", "content+url")}
             for e in ENCODERS}
    return res, mde, means


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-rerun", action="store_true", help="score agreement and noise only")
    args = ap.parse_args()
    key = {r["item"]: r for r in csv.DictReader(open(os.path.join(VDIR, "key.csv"), encoding="utf-8"))}
    A = read_sheet(os.path.join(VDIR, "returned", "cham_trang_A.xlsx"))
    B = read_sheet(os.path.join(VDIR, "returned", "cham_trang_B.xlsx"))
    for name, s in (("A", A), ("B", B)):
        missing = [i for i in key if s.get(i, ("", ""))[0] not in VERDICTS]
        if missing:
            raise SystemExit(f"sheet {name}: {len(missing)} items without a valid verdict, "
                             f"e.g. {missing[:5]}")
    items = sorted(key)
    va = [A[i][0] for i in items]
    vb = [B[i][0] for i in items]
    ca, cb = [collapse(v) for v in va], [collapse(v) for v in vb]
    agree4 = kappa(va, vb, VERDICTS)
    agree3 = kappa(ca, cb, ("abuse", "legitimate", "unsure"))

    cons = {}
    for i, x, y in zip(items, ca, cb):
        cons[i] = x if (x == y and x in ("abuse", "legitimate")) else "unresolved"
    arm = {i: key[i]["in_subset_arm"] for i in items}
    ph_items = [i for i in items if arm[i] == "phishing"]
    be_items = [i for i in items if arm[i] == "benign"]
    ph_res = [i for i in ph_items if cons[i] != "unresolved"]
    ph_legit = sum(cons[i] == "legitimate" for i in ph_res)
    be_res = [i for i in be_items if cons[i] != "unresolved"]
    be_abuse = sum(cons[i] == "abuse" for i in be_res)
    noise_ph = wilson(ph_legit, len(ph_res))
    noise_be = wilson(be_abuse, len(be_res))
    comp = {v: sum(1 for i in ph_items if A[i][0] == v and B[i][0] == v) for v in VERDICTS}
    evid = {e: sum(1 for i in items for s in (A, B) if s[i][1] == e)
            for e in ("capture", "archive", "registry", "name", "none")}

    out = {"n_items": len(items), "agree4": agree4, "agree3": agree3,
           "phishing_arm": {"n": len(ph_items), "resolved": len(ph_res), "legitimate": ph_legit,
                            "noise": noise_ph, "both_same_verdict": comp},
           "benign_arm_rated": {"n": len(be_items), "resolved": len(be_res), "abuse": be_abuse,
                                "noise": noise_be},
           "evidence_votes": evid}

    gen = (f"Two native speakers rated the ${len(ph_items)}$ phishing and ${len(be_items)}$ benign "
           f"pages of the subset that show content (the non-content pages of Section~\\ref{{ssec:indataset}} "
           f"were not rated), blind to label and feed, under the corpus audit's codebook. They agreed on "
           f"${100 * agree3[0]:.1f}\\%$ of the ${len(items)}$ items on the collapsed verdict (abuse, "
           f"legitimate or unsure), with Cohen's $\\kappa={agree3[1]:.2f}$ and PABAK ${agree3[2]:.2f}$. Both "
           f"raters resolved ${len(ph_res)}$ of the ${len(ph_items)}$ phishing pages the same way, and "
           f"${ph_legit}$ of those were legitimate operators (${100 * noise_ph[0]:.1f}\\%$, Wilson "
           f"$95\\%$ CI ${100 * noise_ph[1]:.1f}$--${100 * noise_ph[2]:.1f}\\%$). In the benign "
           f"arm, ${be_abuse}$ of ${len(be_res)}$ resolved pages were abusive")

    if not args.no_rerun:
        import pandas as pd
        from train_content_fusion import load_pages
        from make_p3_relabel_sheets import subset_indices, dead_rows
        from make_p3_assets import ENC_LABEL, SWEEP_METRICS, SWEEP_DIRECTIONS
        path = os.path.join(ROOT, "data", "interim", "content_manifest_vi.csv")
        man = pd.read_csv(path)
        rows = load_pages(path)
        if len(rows) != len(man):
            raise SystemExit("load_pages skipped rows: manifest positions no longer index rows")
        _ph_sub, be_sub = subset_indices(man)
        row_of = {i: int(key[i]["manifest_row"]) for i in items}
        ver_ph = [row_of[i] for i in ph_items if cons[i] == "abuse"]
        drop_be = {row_of[i] for i in be_items if "abuse" in (ca[items.index(i)], cb[items.index(i)])}
        dead = dead_rows(man, be_sub)   # amendment 6: unrated non-content benign pages stay out
        pool = [i for i in be_sub if i not in drop_be and i not in dead]
        ver_be = random.Random(42).sample(pool, len(ver_ph))
        sub = [rows[i] for i in ver_ph] + [rows[i] for i in ver_be]
        res, mde, means = rerun(sub)
        import json
        main_res = json.load(open(os.path.join(ROOT, "data", "processed", "p3", "results",
                                               "encoder_sweep.json"), encoding="utf-8"))
        changed = []
        for e in ENCODERS:
            for d, _ in SWEEP_DIRECTIONS:
                for m, _l, sign in SWEEP_METRICS:
                    old = main_res[e][f"{d}|{m}"]
                    new_ = res[(e, d, m)]
                    was = sign * old["mean"] > 0 and old["q"] < 0.05
                    now = sign * new_["mean"] > 0 and new_["q"] < 0.05
                    flip = (old["mean"] > 0) != (new_["mean"] > 0) and was
                    if was != now or flip:
                        changed.append({"encoder": e, "direction": d, "metric": m,
                                        "was_resolved": was, "now_resolved": now,
                                        "sign_reversed": flip, "mean": new_["mean"],
                                        "q": new_["q"]})
        out["rerun"] = {"n_per_class": len(ver_ph),
                        "results": {"|".join(k): v for k, v in res.items()},
                        "mde": {"|".join(k): v for k, v in mde.items()},
                        "f1_means": means, "changed_verdicts": changed}

        def cell(e, d):
            t = res[(e, d, "F1")]
            body = f"{t['mean']:+.3f}"
            bold = t["mean"] > 0 and t["q"] < 0.05
            return ("\\textbf{" + body + "}" if bold else body) + f" & {{\\scriptsize {_p(t['q'], 'q')}}}"
        body = "\n".join(
            f"{ENC_LABEL.get(e, e)} & {means[e]['url']:.3f} & {means[e]['content']:.3f} & "
            f"{means[e]['content+url']:.3f} & {cell(e, 'cu')} & {cell(e, 'uc')} & "
            f"{mde[(e, 'cu')]:.3f} & {mde[(e, 'uc')]:.3f} \\\\" for e in ENCODERS)
        tex = f"""\\begin{{table}}[t]
\\caption{{The main contrasts on the relabelled subset: ${len(ver_ph)}$ phishing pages both raters
confirmed as abusive and as many benign pages, same protocol as the article (20 splits,
standardised linear head). F1 means, paired F1 differences with BH $q$ per column, and the MDE.}}
\\label{{tab:relabel}}
\\begin{{tabular}}{{l ccc rl rl cc}}
\\toprule
 & \\multicolumn{{3}}{{c}}{{F1}} & \\multicolumn{{2}}{{c}}{{$+$content to URL}} & \\multicolumn{{2}}{{c}}{{$+$URL to content}} & \\multicolumn{{2}}{{c}}{{MDE}} \\\\
Encoder & URL & content & fusion & \\multicolumn{{2}}{{c}}{{$\\Delta$F1}} & \\multicolumn{{2}}{{c}}{{$\\Delta$F1}} & c.\\ to U. & U.\\ to c. \\\\
\\midrule
{body}
\\bottomrule
\\end{{tabular}}
\\end{{table}}"""
        write_generated(os.path.join(SEC, "tab_relabel.tex"), tex, f"(n={len(ver_ph)} per class)")
        n_changed = len(changed)
        gen += (f". On the ${len(ver_ph)}{{+}}{len(ver_ph)}$ pages whose phishing label both raters "
                f"confirmed, the article's protocol "
                + ("returns the same resolved and unresolved verdicts on all $40$ comparisons"
                   if not n_changed else
                   f"changes the verdict of ${n_changed}$ of the $40$ comparisons ("
                   + ", ".join(f"{ENC_LABEL.get(c['encoder'], c['encoder'])} "
                               f"{'content to URL' if c['direction'] == 'cu' else 'URL to content'} "
                               f"{MLAB[c['metric']]} {'now resolved' if c['now_resolved'] else 'no longer resolved'}"
                               + (", sign reversed" if c["sign_reversed"] else "")
                               for c in changed) + ")")
                + " (Supplementary Table~\\SuppRef{tab:relabel})")
    write_results(os.path.join(ROOT, "data", "processed", "p3", "results", "relabel.json"), out)
    write_generated(os.path.join(SEC, "gen_relabel.tex"), gen + "\n")
    print(gen)


if __name__ == "__main__":
    main()
