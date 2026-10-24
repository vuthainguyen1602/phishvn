#!/usr/bin/env python3
"""
make_p3_relabel_sheets.py — build the blinded relabelling package of the P3 page subset.

Plan: papers/P3_multimodal/REVISION2/RELABEL_PLAN.md (committed before this script ran). The 393 content pages of the balanced subset
(amendments 4 and 6: all 414 minus the 21 non-content pages), shuffled, renamed R001-R393. Each rater gets an .xlsx to fill and an offline viewer with the domain, the first 1,500
characters of visible text and the capture's screenshot. The key stays in data/docs/verify_p3/.

RUN:  python scripts/make_p3_relabel_sheets.py
OUT:  data/docs/verify_p3/key.csv                       (withheld from raters)
      data/docs/verify_p3/package/  cham_trang_A.xlsx, cham_trang_B.xlsx, xem_trang.html,
                                    luyen_tap.html, shots/*.jpg, HUONG_DAN.md
"""
import csv
import html
import os
import random
import shutil
import sys

import pandas as pd

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(_HERE))
try:
    from _path import ROOT, add_script_dirs
    add_script_dirs()
except ImportError:  # flat public-mirror layout
    ROOT = os.path.dirname(_HERE)
from train_html_baseline import html_to_text

SEED = 20261017
N_BENIGN = 207   # amendment 4 (2026-10-17): every benign page of the subset is rated
OUT = os.path.join(ROOT, "data", "docs", "verify_p3")
PKG = os.path.join(OUT, "package")
MANIFEST = os.path.join(ROOT, "data", "interim", "content_manifest_vi.csv")
VERDICTS = ["phishing", "scam", "legitimate", "unsure"]
EVIDENCE = ["capture", "archive", "registry", "name", "none"]


def subset_indices(m):
    """The balanced subset exactly as make_p3_assets._balanced_rows draws it: load_pages keeps
    every manifest row (skipped=0), so a row's position in the manifest is its position in the
    row list, and random.Random(42).sample depends only on the population and k."""
    y = (m["label"].astype(str).str.lower() == "phishing").tolist()
    ph = [i for i, v in enumerate(y) if v]
    be = [i for i, v in enumerate(y) if not v]
    rng = random.Random(42)
    k = min(len(ph), len(be))
    return rng.sample(ph, k), rng.sample(be, k)


def dead_rows(m, idx):
    """Manifest rows among idx whose first 600 visible characters match the Vietnamese
    non-content pattern of relabelling amendment 1 (p3_dead_page_check.VI_DEAD). Amendment 6:
    these pages are neither rated nor kept in the verified subset."""
    from p3_dead_page_check import VI_DEAD
    from vn_filter import visible_text
    out = set()
    for i in idx:
        t = visible_text(open(os.path.join(ROOT, m.iloc[i]["dom_file"]), encoding="utf-8",
                              errors="ignore").read())
        if VI_DEAD.search(t[:600]):
            out.add(i)
    return out


def shot_jpg(src, dst, width=1000):
    from PIL import Image
    im = Image.open(src).convert("RGB")
    if im.width > width:
        im = im.resize((width, int(im.height * width / im.width)))
    im.save(dst, "JPEG", quality=70)


def page_text(path, n=1500):
    raw = open(path, encoding="utf-8", errors="ignore").read()
    t = " ".join(html_to_text(raw).split())
    return t[:n] + (" …" if len(t) > n else "")


def viewer(items, title, path, note):
    parts = [f"""<!doctype html><html lang="vi"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title>
<style>body{{font-family:system-ui,sans-serif;max-width:1100px;margin:0 auto;padding:16px;
background:#fafafa;color:#222}} .it{{background:#fff;border:1px solid #ddd;border-radius:8px;
padding:16px;margin:0 0 24px}} h2{{margin:0 0 4px}} .dom{{font-family:monospace;font-size:1.1em;
color:#a40}} img{{max-width:100%;border:1px solid #ccc;margin:8px 0}} .txt{{white-space:pre-wrap;
font-size:.92em;background:#f4f4f4;padding:8px;border-radius:4px}} nav a{{margin-right:6px}}</style>
</head><body><h1>{title}</h1><p>{note}</p><nav>"""]
    parts.append(" ".join(f'<a href="#{it["item"]}">{it["item"]}</a>' for it in items))
    parts.append("</nav>")
    for it in items:
        img = (f'<img src="shots/{it["item"]}.jpg" alt="{it["item"]}" loading="lazy">'
               if it["shot"] else "<p><em>(không có ảnh chụp, chấm theo nội dung chữ)</em></p>")
        parts.append(f'<div class="it" id="{it["item"]}"><h2>{it["item"]}</h2>'
                     f'<div class="dom">{html.escape(it["domain"])}</div>{img}'
                     f'<div class="txt">{html.escape(it["text"])}</div></div>')
    parts.append("</body></html>")
    open(path, "w", encoding="utf-8").write("\n".join(parts))


def sheet(items, path, rater):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "cham"
    head = ["item", "domain", "verdict", "evidence", "brand_impersonated", "note"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for it in items:
        ws.append([it["item"], it["domain"], None, None, None, None])
    dv1 = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(EVIDENCE) + '"', allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"C2:C{len(items) + 1}")
    dv2.add(f"D2:D{len(items) + 1}")
    for col, w in zip("ABCDEF", (8, 42, 14, 12, 24, 40)):
        ws.column_dimensions[col].width = w
    # Amendment 2 (2026-10-17): the evidence column is not collected. It stays in place, hidden,
    # because the locked scoring script reads the sheet's columns by position.
    ws.column_dimensions["D"].hidden = True
    ws.freeze_panes = "C2"
    for row in ws.iter_rows(min_row=2, max_col=2):
        for c in row:
            c.alignment = Alignment(vertical="top")
    wb.properties.creator = f"rater {rater}"
    wb.save(path)


def main():
    m = pd.read_csv(MANIFEST)
    ph, be = subset_indices(m)
    rng = random.Random(SEED)
    be_rated = rng.sample(be, N_BENIGN)
    chosen = ph + be_rated
    rng.shuffle(chosen)
    dead = dead_rows(m, chosen)          # amendment 6: the 21 non-content pages are not rated
    chosen = [i for i in chosen if i not in dead]
    if os.path.exists(PKG):
        shutil.rmtree(PKG)
    os.makedirs(os.path.join(PKG, "shots"))
    items, key = [], []
    for n, i in enumerate(chosen, 1):
        r = m.iloc[i]
        item = f"R{n:03d}"
        shot = str(r.get("shot_file") or "")
        has_shot = bool(shot) and shot != "nan" and os.path.exists(os.path.join(ROOT, shot))
        if has_shot:
            shot_jpg(os.path.join(ROOT, shot), os.path.join(PKG, "shots", f"{item}.jpg"))
        items.append({"item": item, "domain": str(r["domain"]), "shot": has_shot,
                      "text": page_text(os.path.join(ROOT, r["dom_file"]))})
        key.append({"item": item, "manifest_row": i, "domain": r["domain"], "label": r["label"],
                    "provenance": r.get("provenance", ""), "tier": r.get("tier", ""),
                    "in_subset_arm": "phishing" if i in ph else "benign"})
    with open(os.path.join(OUT, "key.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(key[0]))
        w.writeheader()
        w.writerows(key)
    viewer(items, f"Chấm nhãn trang ({len(items)} trang)", os.path.join(PKG, "xem_trang.html"),
           "Mỗi trang: tên miền, ảnh chụp lúc thu thập, 1.500 ký tự chữ đầu tiên. Ghi kết quả vào "
           "file cham_trang_*.xlsx của mình theo đúng mã R001–R393.")
    for rater in ("A", "B"):
        sheet(items, os.path.join(PKG, f"cham_trang_{rater}.xlsx"), rater)

    # Five practice pages from outside the sample: gated Vietnamese captures dropped by the
    # off-corpus cut (never in the subset, so practising on them leaks nothing).
    import build_content_manifest as b
    from vn_filter import is_vietnamese_text, visible_text
    seen = set(m["domain"].astype(str).str.lower())
    pool = []
    det = os.path.join(ROOT, "data", "raw", "urlscan_brands", "detections.csv")
    for r in csv.DictReader(open(det, encoding="utf-8")):
        f = os.path.join(ROOT, (r.get("dom_file") or "").strip())
        d = (r.get("domain") or "").strip().lower()
        if r.get("found") != "1" or d in seen or not os.path.isfile(f):
            continue
        txt = visible_text(open(f, encoding="utf-8", errors="ignore").read())
        if len(txt) < 40 or b.PARKED.search(txt[:400]) or not is_vietnamese_text(txt):
            continue
        pool.append((d, f, (r.get("shot_file") or "").strip()))
    practice = random.Random(SEED).sample(pool, 5)
    pitems = []
    for n, (d, f, s) in enumerate(practice, 1):
        item = f"P{n}"
        ok = bool(s) and os.path.exists(os.path.join(ROOT, s))
        if ok:
            shot_jpg(os.path.join(ROOT, s), os.path.join(PKG, "shots", f"{item}.jpg"))
        pitems.append({"item": item, "domain": d, "shot": ok, "text": page_text(f)})
    viewer(pitems, "Luyện tập (5 trang ngoài mẫu)", os.path.join(PKG, "luyen_tap.html"),
           "Hai người chấm riêng 5 trang này, rồi so cách hiểu bốn nhãn với nhau trước khi bắt đầu "
           "393 trang chính. Kết quả luyện tập không được tính.")
    shutil.copy(os.path.join(ROOT, "papers", "P3_multimodal", "REVISION2", "HUONG_DAN_CHAM_TRANG.md"),
                os.path.join(PKG, "HUONG_DAN.md"))
    n_ph = sum(1 for k in key if k["in_subset_arm"] == "phishing")
    print(f"[+] {len(items)} items ({n_ph} phishing, {len(items) - n_ph} benign), "
          f"{sum(it['shot'] for it in items)} with screenshots -> {PKG}")


if __name__ == "__main__":
    main()
